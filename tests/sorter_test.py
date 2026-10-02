"""The SORT NEW DOWNLOADS button, end to end, against a FAKE downloads
folder built in tests\\_work. Your real E:\\Torrents is never touched —
the test asserts that at the end."""
import sys, os, json, time, shutil, tempfile, subprocess, hashlib
sys.path.insert(0, r"X:\NextUp"); os.chdir(r"X:\NextUp"); sys.argv = ["nextup.py"]
import nextup as n
here = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.join(here, "_work")      # scratch output, never committed
os.makedirs(WORK, exist_ok=True)
sys.path.insert(0, here)
from devserve import isolate
isolate(n, os.path.join(tempfile.gettempdir(), "nextup_sortertest"))

results = []
def check(name, ok, detail=""):
    ok = bool(ok); results.append(ok)
    print("%s  %s%s" % ("PASS" if ok else "FAIL", name, ("  -- %s" % (detail,)) if detail != "" else ""))

REAL = r"E:\Torrents"
real_before = sorted(os.listdir(REAL))

SAND = os.path.join(WORK, "fake_downloads")
if os.path.isdir(SAND):
    shutil.rmtree(SAND)
for sub in ("MOVIES", "TV", "_EXTRAS_QUARANTINE"):
    os.makedirs(os.path.join(SAND, sub))
shutil.copy2(os.path.join(REAL, "process_new_movies.py"), SAND)

# a spreadsheet of the same shape as the real one, so the film path can run
from openpyxl import Workbook
wb = Workbook(); ws = wb.active
ws["A1"], ws["B1"], ws["C1"] = "#", "Title", "Year"
ws["A2"] = "0 films"
wb.save(os.path.join(SAND, "MOVIES", "My Film Collection.xlsx"))

def fake_video(path, mb=1):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(os.urandom(1024) * mb)

# one film, one show that exists in the library already (Futurama), one still downloading
fake_video(os.path.join(SAND, "The.Testing.Film.2019.1080p.BluRay.x265-FAKE", "The.Testing.Film.2019.1080p.mkv"))
fake_video(os.path.join(SAND, "The.Testing.Film.2019.1080p.BluRay.x265-FAKE", "Extras", "Behind.The.Scenes.mkv"))
open(os.path.join(SAND, "The.Testing.Film.2019.1080p.BluRay.x265-FAKE", "rarbg.txt"), "w").write("junk")
for ep in (1, 2, 3):
    fake_video(os.path.join(SAND, "Futurama.S05.1080p.WEBRip.x265-FAKE",
                            "Futurama.S05E%02d.1080p.WEBRip.x265-FAKE.mkv" % ep))
fake_video(os.path.join(SAND, "Something.Still.Downloading.2024.mkv.part"))
# and two that arrived loose, with no folder: a film with its subtitle, an episode
fake_video(os.path.join(SAND, "Loose.Film.2021.1080p.WEBRip.x265-FAKE.mkv"))
open(os.path.join(SAND, "Loose.Film.2021.1080p.WEBRip.x265-FAKE.en.srt"), "w").write("1\n")
fake_video(os.path.join(SAND, "Futurama.S05E04.1080p.WEBRip.x265-FAKE.mkv"))
os.utime(SAND, None)
# the script ignores anything touched in the last 5 minutes
old = time.time() - 3600
for root, dirs, files in os.walk(SAND):
    for f in files + dirs:
        try: os.utime(os.path.join(root, f), (old, old))
        except OSError: pass

n.CONFIG["downloads_folder"] = SAND
print("sandbox:", SAND)

# --- what the button does first: look, change nothing ------------------------
state_before = os.path.isfile(os.path.join(SAND, ".process_new_movies_state.json"))
prev = n.sorter_preview()
check("preview finds all four ready items, loose ones included", len(prev["items"]) == 4,
      [i["name"] + " -> " + i["dest"] for i in prev["items"]])
check("and routes them correctly",
      any("MOVIES/The Testing Film (2019)" in i["dest"] for i in prev["items"])
      and any("MOVIES/Loose Film (2021)" in i["dest"] for i in prev["items"])
      and sum("TV/Futurama" in i["dest"] for i in prev["items"]) == 2,
      [i["dest"] for i in prev["items"]])
check("loose videos listed as waiting, the part-file not",
      "Loose.Film.2021.1080p.WEBRip.x265-FAKE.mkv" in prev["waiting"]
      and not any(w.endswith(".part") for w in prev["waiting"]), prev["waiting"])
check("preview moved nothing", os.path.isdir(os.path.join(SAND, "Futurama.S05.1080p.WEBRip.x265-FAKE"))
      and os.path.isfile(os.path.join(SAND, "Loose.Film.2021.1080p.WEBRip.x265-FAKE.mkv")))
check("preview left the grace-counter file as it found it",
      os.path.isfile(os.path.join(SAND, ".process_new_movies_state.json")) == state_before)

# --- the real run -------------------------------------------------------------
n.sorter_worker()
check("no error from the run", not SORT_ERR if (SORT_ERR := n.SORT["error"]) else True, n.SORT["error"])
film = os.path.join(SAND, "MOVIES", "The Testing Film (2019)", "The Testing Film (2019).mkv")
check("film moved and renamed", os.path.isfile(film), film.replace(SAND, "…"))
check("junk file stripped", not os.path.exists(os.path.join(SAND, "The.Testing.Film.2019.1080p.BluRay.x265-FAKE")))
quar = os.path.join(SAND, "_EXTRAS_QUARANTINE", "MOVIES", "The Testing Film (2019)", "Extras")
check("extras quarantined", os.path.isdir(quar) and os.listdir(quar), quar.replace(SAND, "…"))
tv_root = os.path.join(SAND, "TV")
shows = os.listdir(tv_root)
check("show folder created", shows == ["Futurama"], shows)
eps = []
for root, _d, files in os.walk(tv_root):
    eps += [f for f in files]
check("four episodes moved, the loose one too", len(eps) == 4, eps)
loose = os.path.join(SAND, "MOVIES", "Loose Film (2021)")
check("loose film moved and renamed",
      os.path.isfile(os.path.join(loose, "Loose Film (2021).mkv")), loose.replace(SAND, "…"))
check("its subtitle went with it",
      os.path.isfile(os.path.join(loose, "Loose.Film.2021.1080p.WEBRip.x265-FAKE.en.srt")),
      os.listdir(loose) if os.path.isdir(loose) else "no folder")
check("no loose leftovers in the root",
      not [f for f in os.listdir(SAND) if f.startswith(("Loose.", "Futurama."))],
      [f for f in os.listdir(SAND) if f.startswith(("Loose.", "Futurama."))])
check("episodes named from TMDB, not left as 'S05E01 - Futurama'",
      all(" - Futurama" in e and not e.startswith("S05E01 - ") for e in eps), sorted(eps))
check("the app reported what it named", any("Futurama" in r for r in n.SORT["renamed"]), n.SORT["renamed"])
wb2 = __import__("openpyxl").load_workbook(os.path.join(SAND, "MOVIES", "My Film Collection.xlsx"))
films = [wb2.active.cell(r, 2).value for r in (3, 4)]
check("both films added to the spreadsheet", films == ["Loose Film", "The Testing Film"], films)
check("part-file left alone", os.path.isfile(os.path.join(SAND, "Something.Still.Downloading.2024.mkv.part")))

# --- nothing new: the honest answer -------------------------------------------
again = n.sorter_preview()
check("a second look finds nothing ready", not again["items"], again["items"])

# --- the real folder, untouched ------------------------------------------------
check("E:\\Torrents untouched", sorted(os.listdir(REAL)) == real_before,
      set(sorted(os.listdir(REAL))) ^ set(real_before))
shutil.rmtree(SAND, ignore_errors=True)
print("\n%d/%d passed" % (sum(results), len(results)))
