# -*- coding: utf-8 -*-
"""Duplicates from a sort: the better copy wins, nothing is deleted, and the
losing copies wait for you to say so. Runs a COPY of your installed sorter
against a fake downloads folder of real (tiny) videos, through Next Up's own
sorter code. The real E:\\Torrents is checked untouched at the end, and the
Recycle Bin is replaced with a scratch folder."""
import sys, os, json, time, shutil, tempfile, subprocess, hashlib, sqlite3

here = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.join(here, "_work")      # scratch output, never committed
os.makedirs(WORK, exist_ok=True)
sys.path.insert(0, r"X:\NextUp"); os.chdir(r"X:\NextUp"); sys.argv = ["nextup.py"]
sys.path.insert(0, here)
import nextup as n
from devserve import isolate
isolate(n, os.path.join(tempfile.gettempdir(), "nextup_dupestest"))

results = []
def check(name, ok, detail=""):
    ok = bool(ok); results.append(ok)
    print("%s  %s%s" % ("PASS" if ok else "FAIL", name, ("  -- %s" % (detail,)) if detail != "" else ""))

REAL = r"E:\Torrents"
def fingerprint():
    out = {}
    for name in os.listdir(REAL):
        p = os.path.join(REAL, name)
        out[name] = (os.path.getsize(p), os.path.getmtime(p)) if os.path.isfile(p) else "dir"
    return out
real_before = fingerprint()

SAND = os.path.join(WORK, "fake_dupes")
if os.path.isdir(SAND):
    shutil.rmtree(SAND)
os.makedirs(SAND)
shutil.copy2(r"E:\Torrents\process_new_movies.py", SAND)   # a copy; the original is never run
for sub in ("MOVIES", "TV", "_EXTRAS_QUARANTINE"):
    os.makedirs(os.path.join(SAND, sub))
from openpyxl import Workbook
wb = Workbook(); ws = wb.active
ws["A1"], ws["B1"], ws["C1"] = "#", "Title", "Year"; ws["A2"] = "0 films"
wb.save(os.path.join(SAND, "MOVIES", "My Film Collection.xlsx"))

def clip(path, w, h, seed):
    """A real two-second video at that size. `seed` varies the picture so
    two clips of one size are still different files."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi",
                    "-i", "testsrc=size=%dx%d:rate=10" % (w, h), "-t", "2",
                    "-vf", "hue=h=%d" % seed, "-c:v", "libx264", "-pix_fmt", "yuv420p", path],
                   check=True)
    return md5(path)

def md5(path):
    with open(path, "rb") as fh:
        return hashlib.md5(fh.read()).hexdigest()

show = os.path.join(SAND, "TV", "Test Show (2020)", "Season 01")
lib = {
    "E01": clip(os.path.join(show, "S01E01 Pilot - Test Show.mkv"), 1920, 1080, 10),
    "E02": clip(os.path.join(show, "S01E02 Second - Test Show.mkv"), 854, 480, 20),
    "E03": clip(os.path.join(show, "S01E03 Third - Test Show.mkv"), 1280, 720, 30),
    "film": clip(os.path.join(SAND, "MOVIES", "Test Film (2019)", "Test Film (2019).mp4"), 854, 480, 40),
    "other": clip(os.path.join(SAND, "MOVIES", "Other Film (2018)", "Other Film (2018).mkv"), 1920, 1080, 50),
}
pack = os.path.join(SAND, "Test.Show.S01.WEB-FAKE")
dl = {
    "E01": clip(os.path.join(pack, "Test.Show.S01E01.mkv"), 854, 480, 60),
    "E02": clip(os.path.join(pack, "Test.Show.S01E02.mkv"), 1920, 1080, 70),
    "E03": clip(os.path.join(pack, "Test.Show.S01E03.mkv"), 1280, 720, 30),
    "E04": clip(os.path.join(pack, "Test.Show.S01E04.mkv"), 1280, 720, 80),
    "film": clip(os.path.join(SAND, "Test.Film.2019.1080p-FAKE", "Test.Film.2019.mkv"), 1920, 1080, 90),
    "other": clip(os.path.join(SAND, "Other.Film.2018.480p-FAKE", "Other.Film.2018.mkv"), 854, 480, 100),
}
old = time.time() - 3600          # the script ignores anything touched in the last 5 minutes
for root, dirs, files in os.walk(SAND):
    for f in files + dirs:
        os.utime(os.path.join(root, f), (old, old))

def all_videos():
    out = set()
    for root, dirs, files in os.walk(SAND):
        if "_bin" in root:
            continue
        for f in files:
            if f.endswith((".mkv", ".mp4")):
                out.add(md5(os.path.join(root, f)))
    return out
videos_before = all_videos()

# the library already knows the film you're about to upgrade, with your rating on it
film_old = os.path.normpath(os.path.join(SAND, "MOVIES", "Test Film (2019)", "Test Film (2019).mp4"))
conn = n.db()
conn.execute("INSERT INTO movies (path,title,year,added,userrating) VALUES (?,?,?,?,?)",
             (film_old, "Test Film", 2019, n.now_stamp(), 8))
conn.commit(); conn.close()

n.CONFIG.update(downloads_folder=SAND, tmdb_key="",
                media_folders=[os.path.join(SAND, "MOVIES"), os.path.join(SAND, "TV")])
n.library_scan_worker = lambda: None        # the fake clips are under the scan's 50MB floor
BIN = os.path.join(SAND, "_bin"); os.makedirs(BIN)
def fake_trash(path):
    shutil.move(path, os.path.join(BIN, os.path.basename(path)))
    return "Recycle Bin"
n.send_to_trash = fake_trash

# --- 1. the preview says so before anything moves -----------------------------
prev = n.sorter_preview()
dests = " | ".join(i["dest"] for i in prev["items"])
check("preview: the show pack is flagged", "YOU ALREADY HAVE 3 OF 4 EPISODES" in dests, dests[:160])
check("preview: films you already have are flagged", dests.count("YOU ALREADY HAVE THIS") == 2)
check("preview moved nothing", all_videos() == videos_before)

# --- 2. the sort --------------------------------------------------------------------
n.SORT["running"] = True
n.sorter_worker()
check("sort ran clean", not n.SORT["error"], n.SORT["error"])
report = "\n".join(n.SORT["report"])

check("NOTHING was deleted: every video is still somewhere",
      all_videos() == videos_before, "%d before, %d after" % (len(videos_before), len(all_videos())))

e = lambda name: os.path.join(show, name)
check("E01: your 1080p kept, the 480p download set aside",
      md5(e("S01E01 Pilot - Test Show.mkv")) == lib["E01"])
check("E02: the 1080p download replaced your 480p, under YOUR filename",
      md5(e("S01E02 Second - Test Show.mkv")) == dl["E02"])
check("E03: same quality, yours kept", md5(e("S01E03 Third - Test Show.mkv")) == lib["E03"])
check("E04: the new episode moved in", os.path.isfile(e("S01E04 - Test Show.mkv")))

film_new = os.path.join(SAND, "MOVIES", "Test Film (2019)", "Test Film (2019).mkv")
check("film: the 1080p .mkv replaced your 480p .mp4",
      os.path.isfile(film_new) and md5(film_new) == dl["film"] and not os.path.exists(film_old))
check("other film: your 1080p kept",
      md5(os.path.join(SAND, "MOVIES", "Other Film (2018)", "Other Film (2018).mkv")) == lib["other"])

conn = n.db()
row = conn.execute("SELECT path, userrating FROM movies WHERE title='Test Film'").fetchall()
conn.close()
check("your rating followed the film to its new path",
      len(row) == 1 and row[0][0] == os.path.normpath(film_new) and row[0][1] == 8, row)

from openpyxl import load_workbook
check("films you already had weren't added to the spreadsheet twice",
      load_workbook(os.path.join(SAND, "MOVIES", "My Film Collection.xlsx")).active["A2"].value == "0 films")
check("the report says what happened", "UPGRADED" in report and "DUPLICATE" in report, report[:200])

# --- 3. the question ------------------------------------------------------------------
aside = n.sorter_aside()
check("five copies are waiting to be asked about", aside["count"] == 5, aside["count"])
check("two of them are your old copies, replaced", aside["upgrades"] == 2)
why = {i["name"]: i["why"] for i in aside["items"]}
check("each says why", all(why.values()) and any("1080p replaced your 480p" in w for w in why.values()), why)
check("the download folders are gone from the top level",
      not [x for x in os.listdir(SAND) if x.endswith("-FAKE")], os.listdir(SAND))
check("the holding folder isn't mistaken for a new download", "_DUPLICATES" not in n.sorter_waiting())

if os.environ.get("STOP_BEFORE_BIN"):
    print("stopped before binning; %d waiting" % n.sorter_aside()["count"]); sys.exit(0)

# --- 4. saying yes ---------------------------------------------------------------------
res = n.sorter_bin_aside()
check("yes: all five went to the bin", res["binned"] == 5 and not res["failed"], res)
check("nothing left waiting", n.sorter_aside()["count"] == 0)
check("the emptied release folders were tidied",
      os.listdir(os.path.join(SAND, "_DUPLICATES")) == ["PENDING.jsonl"], os.listdir(os.path.join(SAND, "_DUPLICATES")))
check("and the library itself is untouched by that",
      md5(e("S01E02 Second - Test Show.mkv")) == dl["E02"] and os.path.isfile(film_new))

# --- 5. the real folder ------------------------------------------------------------------
check("the real E:\\Torrents is exactly as it was", fingerprint() == real_before)

print("\n%d/%d passed" % (sum(results), len(results)))
